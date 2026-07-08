import base64
from io import BytesIO
from urllib.parse import urlparse, unquote

import openpyxl

from auth import ensure_authenticated
from graph.client import graph_get, graph_download, graph_download_url


# ---------------------------------------------------------------------------
# Helpers — URL detection & encoding
# ---------------------------------------------------------------------------

def _is_sharing_link(url: str) -> bool:
    """Kiểm tra URL có phải dạng SharePoint sharing link không.

    Sharing links có dạng: https://tenant.sharepoint.com/:x:/s/Site/FileID
    Ký tự đặc biệt '/:' ở đầu path là dấu hiệu nhận biết.
    """
    return urlparse(url).path.startswith("/:")


def _encode_sharing_url(url: str) -> str:
    """Encode sharing URL sang shareId cho Graph API /shares endpoint.

    Theo tài liệu Microsoft:
      1. Base64url-encode toàn bộ URL
      2. Bỏ ký tự padding '='
      3. Thêm prefix 'u!'
    """
    encoded = base64.urlsafe_b64encode(url.encode()).decode().rstrip("=")
    return f"u!{encoded}"


# ---------------------------------------------------------------------------
# Helpers — URL parsing (direct link)
# ---------------------------------------------------------------------------

def _parse_sharepoint_url(url: str) -> tuple[str, str, str]:
    """Tách URL SharePoint thành (hostname, site_path, file_path).

    Ví dụ:
      https://contoso.sharepoint.com/sites/MySite/Shared Documents/folder/file.xlsx
      → ('contoso.sharepoint.com', '/sites/MySite', 'Shared Documents/folder/file.xlsx')
    """
    parsed = urlparse(url)
    hostname = parsed.netloc
    path = unquote(parsed.path)

    for prefix in ("/sites/", "/teams/"):
        if prefix in path:
            idx = path.index(prefix)
            after_prefix = path[idx + len(prefix):]
            parts = after_prefix.split("/", 1)
            site_path = prefix + parts[0]
            file_path = parts[1] if len(parts) > 1 else ""
            return hostname, site_path, file_path

    # Root site — toàn bộ path là đường dẫn file
    return hostname, "", path.lstrip("/")


# ---------------------------------------------------------------------------
# Helpers — Sharing link resolution
# ---------------------------------------------------------------------------

async def _resolve_sharing_link(url: str, token: str) -> tuple[str, str, int, bytes]:
    """Resolve SharePoint sharing link → (file_name, file_path, file_size, file_bytes).

    Sử dụng Graph API endpoint: GET /shares/{shareId}/driveItem
    """
    share_id = _encode_sharing_url(url)

    # 1. Lấy thông tin driveItem từ sharing link
    item_data = await graph_get(
        f"shares/{share_id}/driveItem",
        token,
        params={
            "$select": "id,name,size,parentReference,@microsoft.graph.downloadUrl"
        },
    )

    file_name = item_data.get("name", "unknown")
    file_size = item_data.get("size", 0)

    # 2. Tạo đường dẫn hiển thị từ parentReference
    parent_ref = item_data.get("parentReference", {})
    parent_path = parent_ref.get("path", "")
    file_path = file_name
    if "root:/" in parent_path:
        rel = parent_path.split("root:/", 1)[1]
        file_path = f"{rel}/{file_name}" if rel else file_name

    # 3. Download file
    # Ưu tiên dùng pre-authenticated downloadUrl (không cần Authorization header)
    download_url = item_data.get("@microsoft.graph.downloadUrl")
    if download_url:
        file_bytes = await graph_download_url(download_url)
    else:
        # Fallback: download qua Graph API với token
        drive_id = parent_ref.get("driveId", "")
        item_id = item_data.get("id", "")
        if not drive_id or not item_id:
            raise ValueError(
                "Không lấy được driveId hoặc itemId từ sharing link. "
                "Kiểm tra quyền truy cập file."
            )
        file_bytes = await graph_download(
            f"drives/{drive_id}/items/{item_id}/content", token
        )

    return file_name, file_path, file_size, file_bytes


# ---------------------------------------------------------------------------
# Helpers — Formatting
# ---------------------------------------------------------------------------

def _format_cell(value) -> str:
    if value is None:
        return ""
    return str(value)


def _format_table(rows: list[list]) -> str:
    if not rows:
        return "(không có dữ liệu)"

    col_widths = [0] * len(rows[0])
    for row in rows:
        for i, cell in enumerate(row):
            if i < len(col_widths):
                col_widths[i] = max(col_widths[i], len(_format_cell(cell)))

    lines = []
    for idx, row in enumerate(rows):
        cells = [_format_cell(cell).ljust(col_widths[i]) for i, cell in enumerate(row)]
        lines.append("| " + " | ".join(cells) + " |")
        if idx == 0:
            sep = "|-" + "-|-".join("-" * w for w in col_widths) + "-|"
            lines.append(sep)

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Main implementation
# ---------------------------------------------------------------------------

async def read_excel_impl(
    sharepoint_url: str,
    sheet_name: str | None = None,
    max_rows: int = 100,
) -> str:
    token = ensure_authenticated()

    # -----------------------------------------------------------------------
    # Nhánh 1: Sharing link (e.g. https://tenant.sharepoint.com/:x:/s/...)
    # -----------------------------------------------------------------------
    if _is_sharing_link(sharepoint_url):
        try:
            file_name, file_path, file_size, file_bytes = await _resolve_sharing_link(
                sharepoint_url, token
            )
        except PermissionError as exc:
            return f"Lỗi quyền truy cập: {exc}"
        except ValueError as exc:
            return f"Lỗi khi resolve sharing link: {exc}"
        except Exception as exc:
            return f"Lỗi không xác định: {exc}"

        # Kiểm tra định dạng sau khi biết tên file thật
        if not file_name.lower().endswith((".xlsx", ".xlsm", ".xlam")):
            return (
                f"Lỗi: Chỉ hỗ trợ file Excel định dạng .xlsx / .xlsm / .xlam. "
                f"File được resolve: {file_name}"
            )

        site_display = sharepoint_url

    # -----------------------------------------------------------------------
    # Nhánh 2: Direct URL (e.g. https://tenant.sharepoint.com/sites/.../file.xlsx)
    # -----------------------------------------------------------------------
    else:
        if not sharepoint_url.lower().endswith((".xlsx", ".xlsm", ".xlam")):
            return (
                "Lỗi: Chỉ hỗ trợ file Excel định dạng .xlsx / .xlsm / .xlam. "
                f"File được cung cấp: {sharepoint_url}"
            )

        hostname, site_path, file_path = _parse_sharepoint_url(sharepoint_url)

        if not hostname:
            return f"Lỗi: Không thể phân tích URL SharePoint: {sharepoint_url}"

        # Lấy site ID
        site_api_path = (
            f"sites/{hostname}:{site_path}" if site_path else f"sites/{hostname}"
        )
        try:
            site_data = await graph_get(site_api_path, token)
        except ValueError:
            return (
                f"Lỗi: Không tìm thấy SharePoint site tại '{hostname}{site_path}'. "
                "Kiểm tra lại URL hoặc quyền truy cập."
            )

        site_id = site_data.get("id", "")
        site_display = site_data.get("webUrl", sharepoint_url)

        # Lấy thông tin file (drive item)
        file_api_path = f"sites/{site_id}/drive/root:/{file_path}"
        try:
            item_data = await graph_get(file_api_path, token)
        except ValueError:
            return (
                f"Lỗi: Không tìm thấy file '{file_path}' trong site. "
                "Kiểm tra lại đường dẫn file trong URL."
            )

        item_id = item_data.get("id", "")
        file_name = item_data.get("name", file_path.split("/")[-1])
        file_size = item_data.get("size", 0)

        # Download nội dung file
        download_path = f"sites/{site_id}/drive/items/{item_id}/content"
        try:
            file_bytes = await graph_download(download_path, token)
        except Exception as exc:
            return f"Lỗi khi tải file: {exc}"

    # -----------------------------------------------------------------------
    # Parse Excel (chung cho cả 2 nhánh)
    # -----------------------------------------------------------------------
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        return f"Lỗi khi đọc file Excel: {exc}"

    sheet_names = wb.sheetnames

    if sheet_name is not None and sheet_name not in sheet_names:
        wb.close()
        return (
            f"Lỗi: Không tìm thấy sheet '{sheet_name}'. "
            f"Các sheet hiện có: {', '.join(sheet_names)}"
        )

    target_sheet = sheet_name or sheet_names[0]
    ws = wb[target_sheet]

    rows_data: list[list] = []
    total_rows = 0
    for row in ws.iter_rows(values_only=True):
        if all(cell is None for cell in row):
            continue
        total_rows += 1
        if len(rows_data) < max_rows + 1:  # +1 để tính cả header
            rows_data.append(list(row))

    wb.close()

    # Format output
    lines = [
        f"SharePoint Excel: {file_path}",
        f"Site: {site_display}",
        f"File: {file_name} ({file_size:,} bytes)",
        "",
        f"Sheets: {', '.join(sheet_names)}",
        "",
        f"=== Sheet: {target_sheet} ===",
    ]

    if total_rows == 0:
        lines.append("(sheet trống)")
    else:
        shown = len(rows_data)
        if total_rows > max_rows + 1:
            lines.append(
                f"Tổng số dòng: {total_rows} "
                f"(hiển thị {shown - 1} dòng dữ liệu, bỏ qua header)"
            )
        lines.append("")
        lines.append(_format_table(rows_data))
        if total_rows > max_rows + 1:
            lines.append(
                f"\n... (còn {total_rows - shown} dòng nữa, tăng max_rows để xem thêm)"
            )

    return "\n".join(lines)
