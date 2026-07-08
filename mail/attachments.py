import base64
from io import BytesIO
from pathlib import Path

import openpyxl
import pypdf
import docx

from auth import ensure_authenticated
from graph.client import graph_get
from config import ATTACHMENT_LIST_SELECT


_TEXT_CONTENT_TYPES = {
    "text/plain",
    "text/csv",
    "text/html",
    "text/xml",
    "text/markdown",
    "application/json",
    "application/xml",
    "application/javascript",
}

_TEXT_EXTENSIONS = {".txt", ".csv", ".json", ".html", ".htm", ".xml", ".md"}


def _format_size(size_bytes: int) -> str:
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    return f"{size_bytes / (1024 * 1024):.1f} MB"


def _format_cell(value) -> str:
    return "" if value is None else str(value)


def _format_excel_table(rows: list[list]) -> str:
    if not rows:
        return "(không có dữ liệu)"
    col_widths = [0] * len(rows[0])
    for row in rows:
        for i, cell in enumerate(row):
            if i < len(col_widths):
                col_widths[i] = max(col_widths[i], len(_format_cell(cell)))
    lines = []
    for idx, row in enumerate(rows):
        cells = [_format_cell(cell).ljust(col_widths[i]) for i, cell in enumerate(row)]
        lines.append("| " + " | ".join(cells) + " |")
        if idx == 0:
            sep = "|-" + "-|-".join("-" * w for w in col_widths) + "-|"
            lines.append(sep)
    return "\n".join(lines)


def _read_excel(raw_bytes: bytes, filename: str) -> str:
    try:
        wb = openpyxl.load_workbook(BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as exc:
        return f"Lỗi khi đọc file Excel: {exc}"

    sheet_names = wb.sheetnames
    parts = [f"File: {filename}", f"Sheets: {', '.join(sheet_names)}", ""]

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        rows_data = []
        for row in ws.iter_rows(values_only=True):
            if all(cell is None for cell in row):
                continue
            rows_data.append(list(row))

        parts.append(f"=== Sheet: {sheet_name} ===")
        parts.append(_format_excel_table(rows_data) if rows_data else "(sheet trống)")
        parts.append("")

    wb.close()
    return "\n".join(parts)


def _read_pdf(raw_bytes: bytes, filename: str) -> str:
    try:
        reader = pypdf.PdfReader(BytesIO(raw_bytes))
        total = len(reader.pages)
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"--- Trang {i + 1}/{total} ---\n{text}")
        if not pages_text:
            return f"File: {filename}\n\nKhông trích xuất được text từ PDF này (có thể là PDF ảnh hoặc được mã hóa)."
        return f"File: {filename} ({total} trang)\n\n" + "\n\n".join(pages_text)
    except Exception as exc:
        return f"Lỗi khi đọc file PDF: {exc}"


def _read_docx(raw_bytes: bytes, filename: str) -> str:
    try:
        document = docx.Document(BytesIO(raw_bytes))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        if not paragraphs:
            return f"File: {filename}\n\n(Tài liệu trống hoặc không có nội dung văn bản)"
        return f"File: {filename}\n\n" + "\n".join(paragraphs)
    except Exception as exc:
        return f"Lỗi khi đọc file Word: {exc}"


async def list_attachments_impl(message_id: str) -> str:
    token = ensure_authenticated()

    data = await graph_get(
        f"me/messages/{message_id}/attachments",
        token,
        params={"$select": ATTACHMENT_LIST_SELECT},
    )

    attachments = data.get("value", [])
    if not attachments:
        return f"Email không có file đính kèm."

    lines = [f"Tìm thấy {len(attachments)} file đính kèm:\n"]
    for i, att in enumerate(attachments, 1):
        odata_type = att.get("@odata.type", "")
        att_kind = odata_type.rsplit(".", 1)[-1] if odata_type else "unknown"
        size_label = _format_size(att.get("size", 0))
        inline_label = " [inline]" if att.get("isInline") else ""

        lines.append(
            f"{i}. {att.get('name', '(không có tên)')}{inline_label}\n"
            f"   Loại MIME  : {att.get('contentType', 'unknown')}\n"
            f"   Kích thước : {size_label}\n"
            f"   Kiểu       : {att_kind}\n"
            f"   ID         : {att.get('id', '')}"
        )

    return "\n\n".join(lines)


async def read_attachment_impl(message_id: str, attachment_id: str) -> str:
    token = ensure_authenticated()

    att = await graph_get(
        f"me/messages/{message_id}/attachments/{attachment_id}",
        token,
    )

    odata_type = att.get("@odata.type", "")
    name = att.get("name", "attachment")
    content_type = att.get("contentType", "application/octet-stream")
    size_label = _format_size(att.get("size", 0))
    header = f"Tên       : {name}\nLoại MIME : {content_type}\nKích thước: {size_label}\n" + "─" * 60

    if "itemAttachment" in odata_type:
        return f"{header}\nĐây là item đính kèm (email hoặc lịch liên kết), không có nội dung file."

    if "referenceAttachment" in odata_type:
        web_url = att.get("sourceUrl", "(không có URL)")
        return f"{header}\nLiên kết đám mây: {web_url}"

    content_bytes_b64 = att.get("contentBytes")
    if not content_bytes_b64:
        return f"{header}\nKhông có nội dung (contentBytes trống)."

    raw_bytes = base64.b64decode(content_bytes_b64)
    ext = Path(name).suffix.lower()
    base_mime = content_type.split(";")[0].strip().lower()

    # Excel
    if ext in {".xlsx", ".xlsm", ".xlam"}:
        return f"{header}\n\n{_read_excel(raw_bytes, name)}"

    # PDF
    if ext == ".pdf" or base_mime == "application/pdf":
        return f"{header}\n\n{_read_pdf(raw_bytes, name)}"

    # Word
    if ext == ".docx":
        return f"{header}\n\n{_read_docx(raw_bytes, name)}"

    # Text-based
    if ext in _TEXT_EXTENSIONS or base_mime in _TEXT_CONTENT_TYPES or base_mime.startswith("text/"):
        try:
            return f"{header}\n\n{raw_bytes.decode('utf-8')}"
        except UnicodeDecodeError:
            pass

    # Image
    if base_mime.startswith("image/"):
        return f"{header}\nKhông thể hiển thị nội dung ảnh dưới dạng văn bản."

    return f"{header}\nĐịnh dạng '{content_type}' chưa được hỗ trợ đọc nội dung trực tiếp."
